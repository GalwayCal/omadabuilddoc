"""Builds the as-built Word document from the collected data (python-docx)."""
from __future__ import annotations

import datetime as dt
import io
import logging
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from . import diagrams as dg

log = logging.getLogger(__name__)

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
HEADER_FILL = "1F4E79"
ZEBRA_FILL = "F2F6FA"

DEVICE_STATUS = {0: "Disconnected", 1: "Connected", 2: "Pending", 3: "Heartbeat missed", 4: "Isolated",
                 14: "Connected", 15: "Connected (wireless)", 20: "Pending", 30: "Heartbeat missed"}
LINK_SPEED = {0: "Auto", 1: "10M", 2: "100M", 3: "1G", 4: "2.5G", 5: "10G", 6: "5G", 7: "25G", 8: "40G", 9: "100G"}
WAN_PROTO = {0: "Static", 1: "DHCP", 2: "PPPoE", 3: "L2TP", 4: "PPTP"}
PORT_TYPE = {0: "WAN", 1: "WAN/LAN", 2: "LAN"}
PMF = {1: "Mandatory", 2: "Capable", 3: "Disabled"}
LAN_PROFILE_TYPE = {0: "All", 1: "Disable"}


class Report:
    def __init__(self, data: dict, *, title: str, customer: str = "", author: str = "",
                 secrets: str = "show", floorplans: str | None = None, logo: str | None = None):
        self.data = data
        self.title = title
        self.customer = customer
        self.author = author
        self.secrets = secrets  # show | mask | omit
        self.floorplans = Path(floorplans) if floorplans else None
        self.logo = logo
        self.doc = Document()
        self._fig = 0
        self._setup_styles()

    # =================================================================== API
    def build(self, out_path: str) -> str:
        self._cover()
        self._toc()
        self._controller()
        self._site_summary()
        for site in self.data.get("sites", []):
            self._site(site)
        self._appendix()
        self._enable_update_fields()
        self.doc.save(out_path)
        return out_path

    # ============================================================ formatting
    def _setup_styles(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.0))
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
        st = d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_after = Pt(4)
        for lvl, size in ((1, 18), (2, 13.5), (3, 11.5)):
            h = d.styles[f"Heading {lvl}"]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.color.rgb = ACCENT
            h.font.bold = True
            rf = h.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), "Calibri")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(4)
        cap = d.styles["Caption"]
        cap.font.size = Pt(8.5)
        cap.font.italic = True
        cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        # header / footer
        hdr = sec.header.paragraphs[0]
        hdr.text = f"{self.title}" + (f" — {self.customer}" if self.customer else "")
        hdr.style.font.size = Pt(8)
        hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hdr.runs:
            r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x70, 0x70, 0x70)
        ftr = sec.footer.paragraphs[0]
        ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = ftr.add_run("CONFIDENTIAL — contains network credentials    |    Page "
                        if self.secrets == "show" else "CONFIDENTIAL    |    Page ")
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x70, 0x70, 0x70)
        self._field(ftr, "PAGE", size=8)
        r = ftr.add_run(" of ")
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x70, 0x70, 0x70)
        self._field(ftr, "NUMPAGES", size=8)
        sec.different_first_page_header_footer = True

    @staticmethod
    def _field(par, code: str, size: float = 10, placeholder: str = "1"):
        run = par.add_run()
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
        for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate")):
            el = OxmlElement(tag)
            if attr:
                el.set(qn("w:fldCharType"), attr)
            else:
                el.set(qn("xml:space"), "preserve")
                el.text = f" {code} "
            run._r.append(el)
        t = OxmlElement("w:t")
        t.text = placeholder
        run._r.append(t)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(end)

    def _enable_update_fields(self):
        settings = self.doc.settings.element
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)

    def _shade(self, cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def table(self, headers: list[str], rows: list[list], widths: list[float] | None = None,
              font: float = 8, key_value: bool = False):
        """widths in cm; they should add up to ~17 (A4 minus margins)."""
        if not rows:
            self.note("None configured / nothing returned by the controller.")
            return None
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        hdr = t.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        trPr.append(rep)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(str(h))
            run.bold, run.font.size = True, Pt(font)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            self._shade(c, HEADER_FILL)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run("" if v is None else str(v))
                run.font.size = Pt(font)
                if key_value and i == 0:
                    run.bold = True
                if ri % 2 == 1:
                    self._shade(cells[i], ZEBRA_FILL)
        if widths:
            total = sum(widths)
            widths = [w * 17.0 / total for w in widths]  # always fill the text width
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
            grid = t._tbl.tblGrid
            for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
                gc.set(qn("w:w"), str(int(Cm(widths[i]).twips)))
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            t._tbl.tblPr.append(layout)
        # tighter cell margins
        tblPr = t._tbl.tblPr
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("left", 70), ("right", 70), ("top", 20), ("bottom", 20)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblPr.append(mar)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv(self, pairs: list[tuple[str, object]]):
        rows = [[k, v] for k, v in pairs if v not in (None, "", [], {})]
        return self.table(["Setting", "Value"], rows, widths=[5.5, 11.5], font=9, key_value=True)

    def note(self, text: str, italic=True):
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.italic = italic
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        return p

    def figure(self, png: bytes | None, caption: str, max_w_in=6.6, max_h_in=8.6):
        if not png:
            self.note("Diagram not available (no data).")
            return
        from PIL import Image  # matplotlib depends on Pillow, so it is present
        im = Image.open(io.BytesIO(png))
        w, h = im.size
        dpi = (im.info.get("dpi") or (200, 200))[0] or 200
        # keep diagrams near their designed scale so text size stays consistent
        width = min(max_w_in, max_h_in * w / h, w / dpi * 0.9)
        self.doc.add_picture(io.BytesIO(png), width=Inches(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        self._fig += 1
        self.doc.add_paragraph(f"Figure {self._fig}: {caption}", style="Caption").alignment = \
            WD_ALIGN_PARAGRAPH.CENTER

    def secret(self, value) -> str:
        if value in (None, ""):
            return ""
        if self.secrets == "show":
            return str(value)
        if self.secrets == "mask":
            return "•" * 8
        return "(omitted)"

    def _hyperlink(self, par, url: str, text: str):
        part = self.doc.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "1F5FA8")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(c)
        rPr.append(u)
        r.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        h.append(r)
        par._p.append(h)

    # =============================================================== sections
    def _cover(self):
        d = self.doc
        for _ in range(5):
            d.add_paragraph()
        if self.logo and Path(self.logo).exists():
            d.add_picture(self.logo, width=Cm(5))
        p = d.add_paragraph()
        r = p.add_run(self.title)
        r.font.size, r.bold, r.font.color.rgb = Pt(30), True, ACCENT
        p = d.add_paragraph()
        r = p.add_run("Network As-Built Documentation")
        r.font.size, r.font.color.rgb = Pt(15), RGBColor(0x55, 0x55, 0x55)
        if self.customer:
            p = d.add_paragraph()
            r = p.add_run(self.customer)
            r.font.size, r.bold = Pt(14), True
        d.add_paragraph()
        ctl = self.data.get("controller", {})
        st = ctl.get("status") or {}
        info = ctl.get("info") or {}
        gen = self.data.get("generated", "")
        try:
            gen_disp = dt.datetime.fromisoformat(gen).strftime("%d %B %Y %H:%M %Z").strip()
        except ValueError:
            gen_disp = gen
        rows = [
            ("Generated", gen_disp),
            ("Prepared by", self.author),
            ("Controller", st.get("name") or ctl.get("url")),
            ("Controller version", st.get("controllerVersion") or info.get("controllerVer")),
            ("Sites documented", len(self.data.get("sites", []))),
            ("Source", "Live data from Omada Open API"),
        ]
        self.kv(rows)
        if self.secrets == "show":
            p = d.add_paragraph()
            r = p.add_run("This document contains Wi-Fi passphrases and other credentials in clear text. "
                          "Store and distribute it accordingly.")
            r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), RGBColor(0xB0, 0x2A, 0x2A)
        d.add_page_break()

    def _toc(self):
        self.doc.add_heading("Contents", 1)
        p = self.doc.add_paragraph()
        run = p.add_run()
        for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate")):
            el = OxmlElement(tag)
            if attr:
                el.set(qn("w:fldCharType"), attr)
            else:
                el.set(qn("xml:space"), "preserve")
                el.text = ' TOC \\o "1-2" \\h \\z \\u '
            run._r.append(el)
        t = OxmlElement("w:t")
        t.text = "Right-click and choose Update Field to build the table of contents."
        run._r.append(t)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(end)
        self.doc.add_page_break()

    def _controller(self):
        self.doc.add_heading("1. Controller", 1)
        ctl = self.data.get("controller", {})
        st, info, gen = ctl.get("status") or {}, ctl.get("info") or {}, ctl.get("general") or {}
        cap = st.get("deviceCapacity") or {}
        up = st.get("upTime")
        self.kv([
            ("Controller name", st.get("name") or gen.get("name")),
            ("Management URL", ctl.get("url")),
            ("Omada ID (omadacId)", ctl.get("omadacId")),
            ("Version", st.get("controllerVersion") or info.get("controllerVer")),
            ("API version", info.get("apiVer")),
            ("Type / category", st.get("category")),
            ("Model", st.get("model")),
            ("MAC address", st.get("macAddress")),
            ("Time zone", st.get("timeZone")),
            ("Uptime", f"{up // 86400} days" if isinstance(up, int) else None),
            ("Inform URL", st.get("informUrl")),
            ("Adopted APs / capacity", f"{cap.get('adoptedApNum')} / {cap.get('apCapacity')}" if cap else None),
            ("Adopted switches / capacity", f"{cap.get('adoptedOswNum')} / {cap.get('oswCapacity')}" if cap else None),
            ("Adopted gateways / capacity", f"{cap.get('adoptedOsgNum')} / {cap.get('osgCapacity')}" if cap else None),
        ])

    def _site_summary(self):
        self.doc.add_heading("2. Site Summary", 1)
        rows = []
        for s in self.data.get("sites", []):
            devs = s.get("devices", [])
            cnt = lambda t: sum(1 for d in devs if (d.get("type") or "").lower() == t)  # noqa: E731
            ssids = sum(len(w.get("ssids", [])) for w in s.get("wlan_groups", []))
            i = s.get("info", {})
            rows.append([s.get("name"), i.get("region"), i.get("timeZone"), cnt("gateway"), cnt("switch"),
                         cnt("ap"), len(s.get("lan_networks", [])), ssids])
        self.table(["Site", "Region", "Time zone", "GW", "Switches", "APs", "Networks", "SSIDs"], rows,
                   widths=[4.2, 1.8, 4.0, 1.1, 1.6, 1.3, 1.7, 1.3], font=8.5)

    # ------------------------------------------------------------------ site
    def _site(self, site: dict):
        d = self.doc
        d.add_page_break()
        d.add_heading(f"Site: {site.get('name')}", 1)
        nets_by_id = {n.get("id"): n for n in site.get("lan_networks", [])}

        # ---- overview & location
        d.add_heading("Site details", 2)
        i = site.get("info", {})
        dst = i.get("dst") or {}
        self.kv([
            ("Site name", i.get("name")), ("Site ID", site.get("siteId")),
            ("Scenario", i.get("scenario")), ("Region", i.get("region")), ("Time zone", i.get("timeZone")),
            ("Address", i.get("address")),
            ("Coordinates", f"{i.get('latitude')}, {i.get('longitude')}" if i.get("latitude") else None),
            ("NTP", ", ".join(i.get("ntpServers") or []) if i.get("ntpEnable") else None),
            ("Daylight saving", {0: "Disabled", 1: "Auto", 2: "Manual"}.get(dst.get("mode")) if dst else None),
            ("Public IP", i.get("sitePublicIp")),
        ])
        self._maps(site)

        # ---- devices
        d.add_heading("Device inventory", 2)
        devs = sorted(site.get("devices", []), key=lambda x: ({"gateway": 0, "switch": 1, "ap": 2}.get(
            (x.get("type") or "").lower(), 9), (x.get("name") or "").lower()))
        self.table(["Name", "Type", "Model", "MAC", "IP", "Serial", "Firmware", "Status"],
                   [[x.get("name"), (x.get("type") or "").upper() if x.get("type") == "ap" else
                     (x.get("type") or "").title(), x.get("model") or x.get("modelName"), x.get("mac"),
                     x.get("ip"), x.get("sn"), (x.get("firmwareVersion") or "").split(" Build")[0],
                     DEVICE_STATUS.get(x.get("status"), x.get("status"))] for x in devs],
                   widths=[2.9, 1.3, 2.5, 2.6, 2.1, 2.3, 1.7, 1.6], font=7.5)

        # ---- topology
        d.add_heading("Physical topology", 2)
        wan = self._wan_rows(site)
        wan_label = ", ".join(r[3].split(" /")[0] for r in wan if r[3])[:40]
        self.figure(dg.topology_png(site, wan_label), f"{site.get('name')} — physical topology")
        up_rows = [[x.get("name"), x.get("uplinkDeviceName"), x.get("uplinkDevicePort"),
                    LINK_SPEED.get(x.get("linkSpeed"), x.get("linkSpeed")),
                    {0: "Auto", 1: "Half", 2: "Full"}.get(x.get("duplex"), x.get("duplex"))]
                   for x in devs if x.get("uplinkDeviceMac")]
        if up_rows:
            d.add_heading("Uplinks", 3)
            self.table(["Device", "Uplink device", "Uplink port", "Speed", "Duplex"], up_rows,
                       widths=[4.5, 4.5, 2.6, 2.7, 2.7], font=8)

        # ---- internet
        d.add_heading("Internet / WAN", 2)
        self.table(["Gateway", "Port", "Type", "IP address", "Gateway / DNS", "Status"], wan,
                   widths=[3.0, 2.2, 1.8, 3.2, 4.8, 2.0], font=8)
        cfg = self._wan_config_rows(site)
        if cfg:
            d.add_heading("WAN connection settings", 3)
            self.table(["Port", "Connection", "VLAN", "Username", "Password / details"], cfg,
                       widths=[3.0, 2.2, 1.5, 4.0, 6.3], font=8)

        # ---- networks
        d.add_heading("Networks & VLANs", 2)
        rows = []
        for n in sorted(site.get("lan_networks", []), key=lambda n: n.get("vlan") or 0):
            dh = n.get("dhcpSettingsVO") or {}
            pools = "; ".join(f"{p.get('ipaddrStart')} – {p.get('ipaddrEnd')}" for p in dh.get("ipRangePool") or [])
            dns = ", ".join(x for x in (dh.get("priDns"), dh.get("sndDns")) if x) or \
                ("Auto" if dh.get("dhcpns") == "auto" else "")
            rows.append([n.get("vlan") if n.get("purpose", 1) != 0 or n.get("vlan") else "",
                         n.get("name"), n.get("gatewaySubnet") or "L2 only",
                         ("On" if dh.get("enable") else "Off") if dh else "", pools, dns,
                         dh.get("leasetime"), n.get("domain"),
                         "Yes" if n.get("igmpSnoopEnable") else ""])
        self.table(["VLAN", "Name", "Gateway / subnet", "DHCP", "DHCP range", "DNS", "Lease (min)",
                    "Domain", "IGMP"], rows, widths=[1.1, 2.7, 2.8, 1.1, 3.6, 2.4, 1.3, 1.1, 1.0], font=7.5)

        # ---- port profiles
        profs = site.get("lan_profiles", [])
        prof_by_id = {p.get("id"): p for p in profs}
        d.add_heading("Switch port profiles", 2)

        def nname(nid):
            n = nets_by_id.get(nid)
            return f"{n.get('name')} ({n.get('vlan')})" if n else nid

        self.table(["Profile", "Native network", "Tagged networks", "Voice", "PoE", "Isolation / STP"],
                   [[p.get("name"), nname(p.get("nativeNetworkId")) or LAN_PROFILE_TYPE.get(p.get("type"), ""),
                     ", ".join(nname(x) for x in p.get("tagNetworkIds") or []) or
                     ("All" if p.get("type") == 0 else ""),
                     nname(p.get("voiceNetworkId")) if p.get("voiceNetworkId") else "",
                     {0: "On", 1: "Off", 2: "Unchanged"}.get(p.get("poe"), ""),
                     " / ".join(filter(None, ["Isolated" if p.get("portIsolationEnable") else "",
                                              "STP" if p.get("spanningTreeEnable") else ""]))]
                    for p in profs], widths=[3.2, 3.2, 5.6, 2.0, 1.2, 1.8], font=7.5)

        # ---- switch ports
        switches = [x for x in devs if (x.get("type") or "").lower() == "switch"]
        if switches:
            d.add_heading("Switch port assignments", 2)
            for sw in switches:
                det = (site.get("device_detail", {}).get(sw.get("mac")) or {}).get("switch") or {}
                ports = det.get("portList") or []
                d.add_heading(f"{sw.get('name')} — {sw.get('model') or ''} ({sw.get('ip') or ''})", 3)
                downlinks = {}
                for x in devs:
                    if x.get("uplinkDeviceMac") == sw.get("mac") and x.get("uplinkDevicePort"):
                        downlinks[str(x["uplinkDevicePort"])] = x.get("name")
                self.table(["Port", "Name", "Profile", "PoE", "Connected device", "State"],
                           [[p.get("port"), p.get("name"),
                             p.get("profileName") or (prof_by_id.get(p.get("profileId")) or {}).get("name"),
                             "On" if p.get("poeMode") == 1 else ("Off" if p.get("poeMode") == 0 else ""),
                             downlinks.get(str(p.get("port")), ""),
                             "LAG" if p.get("lagPort") else ({1: "Enabled", 0: "Disabled"}.get(p.get("status"), ""))]
                            for p in sorted(ports, key=lambda p: p.get("port") or 0)],
                           widths=[1.2, 4.0, 3.6, 1.2, 4.6, 2.4], font=7.5)

        # ---- wireless
        d.add_heading("Wireless networks", 2)
        radius = {r.get("id") or r.get("radiusProfileId"): r.get("name") for r in site.get("radius_profiles", [])}
        ppsk = {r.get("id"): r.get("name") for r in site.get("ppsk_profiles", [])}
        default_vlan = next((n.get("vlan") for n in site.get("lan_networks", []) if n.get("primary")), 1)
        rows = []
        for w in site.get("wlan_groups", []):
            for s in w.get("ssids", []):
                key = ((s.get("pskSetting") or {}).get("securityKey") if s.get("security") == 3 else "")
                if s.get("security") == 2:
                    key = f"RADIUS: {radius.get((s.get('entSetting') or {}).get('radiusProfileId'), 'profile')}"
                elif s.get("security") in (4, 5):
                    key = f"PPSK: {ppsk.get((s.get('ppskSetting') or {}).get('ppskProfileId'), 'profile')}"
                else:
                    key = self.secret(key)
                rows.append([w.get("name"), s.get("name"), dg.band_text(s.get("band")), dg.security_text(s), key,
                             dg.ssid_vlan(s, default_vlan), "No" if s.get("broadcast") is False else "Yes",
                             "Yes" if s.get("guestNetEnable") else "", "Yes" if s.get("enable11r") else "",
                             PMF.get(s.get("pmfMode"), "")])
        self.table(["WLAN group", "SSID", "Bands", "Security", "Passphrase / auth", "VLAN", "Broadcast",
                    "Guest", "802.11r", "PMF"], rows,
                   widths=[2.0, 2.6, 1.5, 2.1, 3.3, 1.0, 1.3, 1.0, 1.1, 1.3], font=7.5)
        if self.secrets == "omit":
            self.note("Passphrases omitted from this edition of the document.")

        aps = [x for x in devs if (x.get("type") or "").lower() == "ap"]
        if aps:
            d.add_heading("Access points", 3)
            wlan_names = {w.get("wlanId"): w.get("name") for w in site.get("wlan_groups", [])}
            topo_nodes = {n.get("mac"): n for n in (site.get("topology") or {}).get("topologyNodes") or []}
            rows = []
            for a in aps:
                det = (site.get("device_detail", {}).get(a.get("mac")) or {}).get("ap") or {}
                tn = topo_nodes.get(a.get("mac"), {})
                ch = " / ".join(str(tn.get(k)) for k in ("channel2g", "channel5g", "channel6g") if tn.get(k))
                wu = det.get("wireless uplink info") or {}
                rows.append([a.get("name"), a.get("model"), a.get("ip"),
                             wlan_names.get(det.get("wlan group id"), det.get("wlan group id") or ""),
                             f"Mesh via {wu.get('name')}" if wu.get("uplinkMac") else
                             f"{a.get('uplinkDeviceName') or ''} p{a.get('uplinkDevicePort') or ''}".strip(" p"),
                             ch])
            self.table(["AP", "Model", "IP", "WLAN group", "Uplink", "Channels (2.4/5/6)"], rows,
                       widths=[3.2, 2.6, 2.2, 2.6, 4.0, 2.4], font=7.5)

        d.add_heading("Logical network diagram", 2)
        self.figure(dg.logical_png(site), f"{site.get('name')} — VLANs and SSID mapping")

        # ---- DHCP reservations
        res = site.get("dhcp_reservations", [])
        if res:
            d.add_heading("DHCP reservations", 2)
            self.table(["Network", "IP", "MAC", "Client / description", "Enabled"],
                       [[r.get("netName"), r.get("ip"), r.get("mac"),
                         r.get("clientName") or r.get("description"), "Yes" if r.get("status") else "No"]
                        for r in sorted(res, key=lambda r: tuple(int(o) for o in (r.get("ip") or "0.0.0.0").split(".")
                                                                 if o.isdigit()))],
                       widths=[3.0, 2.8, 3.2, 6.0, 2.0], font=8)

    # ----------------------------------------------------------- site parts
    def _maps(self, site: dict):
        i = site.get("info", {})
        imgs = []
        if self.floorplans and self.floorplans.exists():
            name = (site.get("name") or "").lower()
            for f in sorted(self.floorplans.rglob("*")):
                if f.suffix.lower() in (".png", ".jpg", ".jpeg") and (
                        f.name.lower().startswith(name) or f.parent.name.lower() == name):
                    imgs.append((f.stem, f.read_bytes()))
        import base64
        for m in site.get("maps", []):
            if m.get("_image_b64"):
                imgs.append((m.get("name") or "Map", base64.b64decode(m["_image_b64"])))
        has_geo = i.get("latitude") not in (None, 0) and i.get("longitude") not in (None, 0)
        if not imgs and not has_geo and not site.get("maps"):
            return
        self.doc.add_heading("Maps & location", 2)
        if has_geo:
            p = self.doc.add_paragraph("Site location: ")
            lat, lon = i["latitude"], i["longitude"]
            self._hyperlink(p, f"https://www.openstreetmap.org/?mlat={lat}&mlon={lon}#map=18/{lat}/{lon}",
                            f"{lat}, {lon} (open in OpenStreetMap)")
        if site.get("maps"):
            self.table(["Map", "Details"], [[m.get("name") or m.get("id"),
                                             ", ".join(f"{k}={v}" for k, v in m.items()
                                                       if not k.startswith("_") and isinstance(v, (str, int, float))
                                                       and k not in ("name",))[:180]]
                                            for m in site["maps"]], widths=[4, 13], font=8)
        for name, blob in imgs:
            self.figure(blob, f"{site.get('name')} — {name}")

    def _wan_rows(self, site: dict) -> list:
        rows = []
        for d in site.get("devices", []):
            det = site.get("device_detail", {}).get(d.get("mac")) or {}
            for p in det.get("wan_status") or []:
                if not isinstance(p, dict):
                    continue
                if p.get("mode") not in (0, None) and p.get("type") == 2:
                    continue  # pure LAN port
                v4 = p.get("wanPortIpv4Config") or {}
                ip = v4.get("ip") or p.get("ip") or ""
                if ip and v4.get("netmask"):
                    ip = f"{ip} / {v4['netmask']}"
                gw_dns = "; ".join(filter(None, [f"GW {v4['gateway']}" if v4.get("gateway") else "",
                                                 ", ".join(filter(None, [v4.get("priDns"), v4.get("sndDns")]))]))
                rows.append([d.get("name"), p.get("name") or p.get("portDesc") or p.get("port"),
                             (p.get("proto") or "").upper(), ip, gw_dns,
                             "Up" if p.get("status") == 1 or p.get("internetState") == 1 else "Down"])
        return rows

    def _wan_config_rows(self, site: dict) -> list:
        rows = []
        for w in (site.get("internet") or {}).get("wanPortSettings") or []:
            v4 = w.get("wanPortIpv4Setting") or {}
            proto = v4.get("protoType")
            user, detail = "", ""
            if proto == 0:
                st = v4.get("ipv4Static") or {}
                detail = f"{st.get('ipAddress')}/{st.get('subnetMask')} gw {st.get('defaultGateway')}"
            elif proto in (2, 3, 4):
                cred = v4.get({2: "ipv4Pppoe", 3: "ipv4L2tp", 4: "ipv4Pptp"}[proto]) or {}
                user, detail = cred.get("userName", ""), self.secret(cred.get("password"))
            rows.append([w.get("portName") or w.get("portDescription"), WAN_PROTO.get(proto, proto),
                         v4.get("vlanId") or "", user, detail])
        return rows

    # ------------------------------------------------------------- appendix
    def _appendix(self):
        d = self.doc
        d.add_page_break()
        d.add_heading("Appendix A. Data collection notes", 1)
        self.note("Items the controller did not return. Usually a missing Open API permission on the "
                  "application, a feature not present on this controller build, or no such hardware on "
                  "the site. Raw API output is saved alongside this document.", italic=False)
        rows = [["Controller", e["what"], e["error"][:160]] for e in self.data.get("errors", [])]
        for s in self.data.get("sites", []):
            rows += [[s.get("name"), e["what"], e["error"][:160]] for e in s.get("errors", [])]
        self.table(["Scope", "Item", "Detail"], rows, widths=[3.0, 4.5, 9.5], font=7.5)
